from flask import Blueprint, request, jsonify, current_app
from werkzeug.utils import secure_filename
from src.models.assessment import Document, AssessmentResult, PLOAssessmentSummary, db
from src.models.clo import CLO, CLOPLOMapping, CLOKeyword, ScoringScale
from src.models.plo import PLO
from src.ai_integration import CLOAssessmentEngine, CLOAssessmentConfig, assess_document_for_all_clos, calculate_plo_scores
import os
import uuid
from datetime import datetime
import json
import PyPDF2
import docx
from collections import defaultdict

assessment_bp = Blueprint('assessment', __name__)

ALLOWED_EXTENSIONS = {'txt', 'pdf', 'doc', 'docx', 'rtf'}
UPLOAD_FOLDER = 'uploads'

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def extract_text_from_file(file_path, file_type):
    """Extract text content from uploaded file"""
    try:
        if file_type == 'txt':
            with open(file_path, 'r', encoding='utf-8') as f:
                return f.read()
        
        elif file_type == 'pdf':
            text = ""
            with open(file_path, 'rb') as f:
                pdf_reader = PyPDF2.PdfReader(f)
                for page in pdf_reader.pages:
                    text += page.extract_text() + "\n"
            return text
        
        elif file_type in ['doc', 'docx']:
            doc = docx.Document(file_path)
            text = ""
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            return text
        
        else:
            return ""
    
    except Exception as e:
        current_app.logger.error(f"Error extracting text from {file_path}: {e}")
        return ""

@assessment_bp.route('/documents/upload', methods=['POST'])
def upload_document():
    """Upload document for assessment"""
    try:
        if 'file' not in request.files:
            return jsonify({
                'success': False,
                'error': 'No file provided'
            }), 400
        
        file = request.files['file']
        if file.filename == '':
            return jsonify({
                'success': False,
                'error': 'No file selected'
            }), 400
        
        if not allowed_file(file.filename):
            return jsonify({
                'success': False,
                'error': f'File type not allowed. Allowed types: {", ".join(ALLOWED_EXTENSIONS)}'
            }), 400
        
        # Get additional metadata
        course_code = request.form.get('course_code', '')
        document_type = request.form.get('document_type', 'assignment')
        
        # Create upload directory if it doesn't exist
        upload_dir = os.path.join(current_app.root_path, UPLOAD_FOLDER)
        os.makedirs(upload_dir, exist_ok=True)
        
        # Generate unique filename
        file_extension = file.filename.rsplit('.', 1)[1].lower()
        unique_filename = f"{uuid.uuid4()}.{file_extension}"
        file_path = os.path.join(upload_dir, unique_filename)
        
        # Save file
        file.save(file_path)
        
        # Extract text content
        content_text = extract_text_from_file(file_path, file_extension)
        
        # Create document record
        document = Document(
            filename=unique_filename,
            original_filename=secure_filename(file.filename),
            file_path=file_path,
            file_size=os.path.getsize(file_path),
            file_type=file_extension,
            course_code=course_code,
            document_type=document_type,
            content_text=content_text
        )
        
        db.session.add(document)
        db.session.commit()
        
        return jsonify({
            'success': True,
            'data': document.to_dict(),
            'message': 'Document uploaded successfully'
        }), 201
        
    except Exception as e:
        db.session.rollback()
        return jsonify({
            'success': False,
            'error': str(e)
        }), 500

@assessment_bp.route('/documents', methods=['GET'])
def get_documents():
    """Get all uploaded documents"""
    try:
        course_code = request.args.get('course_code')
        document_type = request.args.get('document_type')
        
        query = Document.query
        if course_code:
            query = query.filter_by(course_code=course_code)
        if document_type:
            query = query.filter_by(document_type=document_type)
        
        documents = query.order_by(Document.upload_date.desc()).all()
        
        return jsonify({
            'success': True,
            'data': [doc.to_dict() for doc in documents],
            'total': len(documents)
        })
        
    except Exception as e:
        return jsonify({
            'success': False,
            'error': str(e)
        }), 500

@assessment_bp.route('/documents/<int:document_id>', methods=['GET'])
def get_document(document_id):
    """Get specific document"""
    try:
        document = Document.query.get_or_404(document_id)
        doc_dict = document.to_dict()
        
        # Add assessment results
        results = AssessmentResult.query.filter_by(document_id=document_id).all()
        doc_dict['assessment_results'] = [result.to_dict() for result in results]
        
        return jsonify({
            'success': True,
            'data': doc_dict
        })
        
    except Exception as e:
        return jsonify({
            'success': False,
            'error': str(e)
        }), 404

@assessment_bp.route('/documents/<int:document_id>/assess', methods=['POST'])
def assess_document(document_id):
    """Assess document against CLOs"""
    try:
        document = Document.query.get_or_404(document_id)
        
        if not document.content_text:
            return jsonify({
                'success': False,
                'error': 'Document has no extractable text content'
            }), 400
        
        # Get CLOs for the course
        clos = CLO.query.filter_by(course_code=document.course_code).all()
        if not clos:
            return jsonify({
                'success': False,
                'error': f'No CLOs found for course {document.course_code}'
            }), 400
        
        # Prepare CLO configurations
        clo_configs = []
        for clo in clos:
            # Get keywords
            keywords_objs = CLOKeyword.query.filter_by(clo_id=clo.id).all()
            keywords = []
            for kw_obj in keywords_objs:
                keywords.extend(kw_obj.keywords.split(','))
            
            # Get scoring config
            scoring = ScoringScale.query.filter_by(clo_id=clo.id).first()
            if not scoring:
                # Use default scoring
                scoring = ScoringScale(
                    method='keyword_frequency',
                    full_score_threshold=3.0,
                    partial_score_floor=0.3,
                    max_score=4.0
                )
            
            config = CLOAssessmentConfig(
                clo_id=clo.id,
                keywords=keywords,
                method=scoring.method,
                full_score_threshold=scoring.full_score_threshold,
                partial_score_floor=scoring.partial_score_floor,
                max_score=scoring.max_score
            )
            clo_configs.append(config)
        
        # Perform assessment
        assessment_results = assess_document_for_all_clos(document.content_text, clo_configs)
        
        # Save results to database
        saved_results = []
        for result in assessment_results:
            assessment_result = AssessmentResult(
                document_id=document_id,
                clo_id=result['clo_id'],
                score=result['score'],
                confidence=result['confidence'],
                assessment_method=result['method']
            )
            assessment_result.set_details(result['details'])
            
            db.session.add(assessment_result)
            saved_results.append(assessment_result)
        
        # Calculate PLO scores
        clo_plo_mappings = {}
        for clo in clos:
            mappings = CLOPLOMapping.query.filter_by(clo_id=clo.id).all()
            clo_plo_mappings[clo.id] = [(m.plo_id, m.weight) for m in mappings]
        
        plo_scores = calculate_plo_scores(assessment_results, clo_plo_mappings)
        
        # Mark document as processed
        document.processed = True
        
        db.session.commit()
        
        return jsonify({
            'success': True,
            'clo_results': [result.to_dict() for result in saved_results],
            'plo_scores': plo_scores,
            'message': 'Document assessed successfully'
        })
        
    except Exception as e:
        db.session.rollback()
        return jsonify({
            'success': False,
            'error': str(e)
        }), 500

@assessment_bp.route('/documents/<int:document_id>/results', methods=['GET'])
def get_assessment_results(document_id):
    """Get assessment results for a document"""
    try:
        document = Document.query.get_or_404(document_id)
        results = AssessmentResult.query.filter_by(document_id=document_id).all()
        
        # Group results by CLO and calculate PLO scores
        clo_results = {}
        for result in results:
            clo_results[result.clo_id] = result.to_dict()
            
            # Add CLO details
            clo = CLO.query.get(result.clo_id)
            if clo:
                clo_results[result.clo_id]['clo'] = {
                    'code': clo.code,
                    'title': clo.title,
                    'course_code': clo.course_code
                }
        
        # Calculate PLO scores
        clo_plo_mappings = {}
        for clo_id in clo_results.keys():
            mappings = CLOPLOMapping.query.filter_by(clo_id=clo_id).all()
            clo_plo_mappings[clo_id] = [(m.plo_id, m.weight) for m in mappings]
        
        assessment_data = [{'clo_id': clo_id, 'score': data['score']} 
                          for clo_id, data in clo_results.items()]
        plo_scores = calculate_plo_scores(assessment_data, clo_plo_mappings)
        
        # Add PLO details
        for plo_id in plo_scores.keys():
            plo = PLO.query.get(plo_id)
            if plo:
                plo_scores[plo_id]['plo'] = {
                    'code': plo.code,
                    'title': plo.title,
                    'category': plo.category
                }
        
        return jsonify({
            'success': True,
            'document': document.to_dict(),
            'clo_results': list(clo_results.values()),
            'plo_scores': plo_scores
        })
        
    except Exception as e:
        return jsonify({
            'success': False,
            'error': str(e)
        }), 500

@assessment_bp.route('/summary/course/<course_code>', methods=['GET'])
def get_course_assessment_summary(course_code):
    """Get assessment summary for a course"""
    try:
        # Get all documents for the course
        documents = Document.query.filter_by(course_code=course_code, processed=True).all()
        
        if not documents:
            return jsonify({
                'success': False,
                'error': f'No processed documents found for course {course_code}'
            }), 404
        
        # Get all assessment results for these documents
        document_ids = [doc.id for doc in documents]
        results = AssessmentResult.query.filter(AssessmentResult.document_id.in_(document_ids)).all()
        
        # Group by CLO
        clo_summary = defaultdict(list)
        for result in results:
            clo_summary[result.clo_id].append(result.score)
        
        # Calculate CLO statistics
        clo_stats = {}
        for clo_id, scores in clo_summary.items():
            clo = CLO.query.get(clo_id)
            if clo:
                avg_score = sum(scores) / len(scores)
                achievement_rate = len([s for s in scores if s >= 2.0]) / len(scores) * 100
                
                clo_stats[clo_id] = {
                    'clo': {
                        'code': clo.code,
                        'title': clo.title
                    },
                    'total_assessments': len(scores),
                    'average_score': round(avg_score, 2),
                    'achievement_rate': round(achievement_rate, 2),
                    'scores': scores
                }
        
        # Calculate PLO summary
        clo_plo_mappings = {}
        for clo_id in clo_stats.keys():
            mappings = CLOPLOMapping.query.filter_by(clo_id=clo_id).all()
            clo_plo_mappings[clo_id] = [(m.plo_id, m.weight) for m in mappings]
        
        # Calculate PLO scores for each document
        plo_document_scores = defaultdict(list)
        for doc in documents:
            doc_results = [r for r in results if r.document_id == doc.id]
            doc_clo_results = [{'clo_id': r.clo_id, 'score': r.score} for r in doc_results]
            
            if doc_clo_results:
                plo_scores = calculate_plo_scores(doc_clo_results, clo_plo_mappings)
                for plo_id, plo_data in plo_scores.items():
                    plo_document_scores[plo_id].append(plo_data['score'])
        
        # Calculate PLO statistics
        plo_stats = {}
        for plo_id, scores in plo_document_scores.items():
            plo = PLO.query.get(plo_id)
            if plo and scores:
                avg_score = sum(scores) / len(scores)
                achievement_rate = len([s for s in scores if s >= 2.0]) / len(scores) * 100
                
                plo_stats[plo_id] = {
                    'plo': {
                        'code': plo.code,
                        'title': plo.title,
                        'category': plo.category
                    },
                    'total_assessments': len(scores),
                    'average_score': round(avg_score, 2),
                    'achievement_rate': round(achievement_rate, 2),
                    'scores': scores
                }
        
        return jsonify({
            'success': True,
            'course_code': course_code,
            'total_documents': len(documents),
            'clo_summary': clo_stats,
            'plo_summary': plo_stats,
            'assessment_period': {
                'start_date': min(doc.upload_date for doc in documents).isoformat(),
                'end_date': max(doc.upload_date for doc in documents).isoformat()
            }
        })
        
    except Exception as e:
        return jsonify({
            'success': False,
            'error': str(e)
        }), 500

@assessment_bp.route('/summary/overall', methods=['GET'])
def get_overall_assessment_summary():
    """Get overall assessment summary across all courses"""
    try:
        # Get all processed documents
        documents = Document.query.filter_by(processed=True).all()
        
        if not documents:
            return jsonify({
                'success': False,
                'error': 'No processed documents found'
            }), 404
        
        # Group by course
        course_summary = defaultdict(list)
        for doc in documents:
            course_summary[doc.course_code].append(doc)
        
        # Calculate summary for each course
        overall_stats = {}
        for course_code, course_docs in course_summary.items():
            document_ids = [doc.id for doc in course_docs]
            results = AssessmentResult.query.filter(AssessmentResult.document_id.in_(document_ids)).all()
            
            if results:
                avg_score = sum(r.score for r in results) / len(results)
                achievement_rate = len([r for r in results if r.score >= 2.0]) / len(results) * 100
                
                overall_stats[course_code] = {
                    'total_documents': len(course_docs),
                    'total_assessments': len(results),
                    'average_score': round(avg_score, 2),
                    'achievement_rate': round(achievement_rate, 2)
                }
        
        return jsonify({
            'success': True,
            'total_courses': len(course_summary),
            'total_documents': len(documents),
            'course_summary': overall_stats
        })
        
    except Exception as e:
        return jsonify({
            'success': False,
            'error': str(e)
        }), 500

@assessment_bp.route('/documents/<int:document_id>', methods=['DELETE'])
def delete_document(document_id):
    """Delete document and all related assessment results"""
    try:
        document = Document.query.get_or_404(document_id)
        
        # Delete file from filesystem
        if os.path.exists(document.file_path):
            os.remove(document.file_path)
        
        # Delete assessment results
        AssessmentResult.query.filter_by(document_id=document_id).delete()
        
        # Delete document record
        db.session.delete(document)
        db.session.commit()
        
        return jsonify({
            'success': True,
            'message': 'Document deleted successfully'
        })
        
    except Exception as e:
        db.session.rollback()
        return jsonify({
            'success': False,
            'error': str(e)
        }), 500

